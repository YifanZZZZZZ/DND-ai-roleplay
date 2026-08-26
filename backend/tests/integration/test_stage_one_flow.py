from pathlib import Path
from uuid import uuid4

from httpx import AsyncClient
from openpyxl import Workbook

from backend.app.core.config import Settings


def build_character_sheet(path: Path, character_name: str) -> None:
    workbook = Workbook()
    main = workbook.active
    main.title = "主要"
    origin = workbook.create_sheet("起源")
    workbook.create_sheet("背包")
    spells = workbook.create_sheet("法术大全")

    main["A1"] = "DND 5E2024 人物卡<悲灵v1.0.0>"
    main["E3"] = character_name
    main["E6"] = "魔契师"
    main["I6"] = "咒剑"
    main["T6"] = "影灵"
    main["AX17"] = "契约魔法"
    main["BC17"] = "角色能够借助契约施展魔法。"
    main["BT30"] = "鸦后的祝福"
    main["BZ30"] = "角色可以短暂穿越阴影。"
    main["BU43"] = "健壮"
    main["BZ43"] = "角色有着异于常人的生命力。"
    main["Q68"] = "魔能爆"
    main["L31"] = "鳞甲"
    main["V31"] = "穿戴时行动较为沉重。"
    skills = {
        32: ("运动", 1),
        34: ("特技", 5),
        35: ("巧手", 2),
        36: ("隐匿", 5),
        38: ("调查", 8),
        39: ("奥秘", 5),
        40: ("历史", 5),
        41: ("自然", 5),
        42: ("宗教", 8),
        44: ("察觉", 0),
        45: ("洞悉", 0),
        46: ("驯兽", 0),
        47: ("医药", 0),
        48: ("求生", 0),
        50: ("游说", 4),
        51: ("欺瞒", 1),
        52: ("威吓", 4),
        53: ("表演", 1),
    }
    for row, (label, value) in skills.items():
        main[f"C{row}"] = label
        main[f"I{row}"] = value

    origin["E5"] = "堕影冥界"
    origin["E6"] = "城市猎人"
    origin["B24"] = "通用语"
    origin["B31"] = "赌具"

    spells["A2"] = "法术名"
    spells["C2"] = "学派"
    spells["F2"] = "施法时间"
    spells["G2"] = "施法距离"
    spells["M2"] = "法术详述"
    spells["A3"] = "魔能爆"
    spells["C3"] = "塑能"
    spells["F3"] = "动作"
    spells["G3"] = "远距离"
    spells["M3"] = "角色释放一束爆裂的魔法能量。"

    workbook.save(path)
    workbook.close()


async def create_character(client: AsyncClient, name: str, max_hp: int = 30) -> dict[str, object]:
    response = await client.post(
        "/api/v1/characters",
        json={"name": name, "roleplayPrompt": f"你是{name}。", "maxHp": max_hp},
    )
    assert response.status_code == 201
    return response.json()


async def configure_sheet(
    client: AsyncClient, tmp_path: Path, character_id: str, character_name: str
) -> None:
    sheet_path = tmp_path / f"{character_name}.xlsx"
    build_character_sheet(sheet_path, character_name)
    with sheet_path.open("rb") as source:
        response = await client.post(
            f"/api/v1/characters/{character_id}/sheet-versions:preview",
            files={
                "file": (
                    sheet_path.name,
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 201
    preview = response.json()
    assert preview["snapshot"]["characterName"] == character_name
    assert preview["snapshot"]["species"] == "影灵"
    assert len(preview["snapshot"]["classFeatures"]) == 1
    assert len(preview["snapshot"]["spells"]) == 1
    assert preview["spellbook"][0]["name"] == "魔能爆"
    assert "角色释放一束爆裂的魔法能量" in preview["spellbook"][0]["summary"]
    assert not any(character.isdigit() for character in preview["spellbook"][0]["summary"])
    assert len(preview["snapshot"]["skills"]) == 18
    assert preview["snapshot"]["skills"]["investigation"] == 8

    activation = await client.post(
        f"/api/v1/characters/{character_id}/sheet-versions/{preview['versionId']}:activate",
        json={
            "spellbook": [
                {
                    "name": "魔能爆",
                    "category": "CANTRIP",
                    "summary": "向视野中的目标释放爆裂魔法能量，命中与结果由 DM 裁定。",
                }
            ]
        },
    )
    assert activation.status_code == 200
    assert activation.json()["spellbook"][0]["summary"].startswith("向视野中的目标")

    skills_response = await client.get(f"/api/v1/characters/{character_id}/skills")
    assert skills_response.status_code == 200
    modifiers = skills_response.json()["modifiers"]
    assert modifiers["investigation"] == 8
    assert modifiers["religion"] == 8
    assert modifiers["perception"] == 0


async def test_spellbook_can_be_manually_updated_after_sheet_activation(
    api_client: AsyncClient, tmp_path: Path
) -> None:
    character = await create_character(api_client, "卡斯珀")
    character_id = str(character["id"])
    await configure_sheet(api_client, tmp_path, character_id, "卡斯珀")

    current = await api_client.get(f"/api/v1/characters/{character_id}/spellbook")
    assert current.status_code == 200
    payload = current.json()
    assert [spell["name"] for spell in payload["spells"]] == ["魔能爆"]

    updated = await api_client.put(
        f"/api/v1/characters/{character_id}/spellbook",
        json={
            "revision": payload["revision"],
            "spells": [
                payload["spells"][0],
                {
                    "name": "护盾术",
                    "category": "PREPARED",
                    "summary": "迅速形成短暂的魔法屏障，是否挡住来袭攻击由 DM 裁定。",
                },
            ],
        },
    )
    assert updated.status_code == 200
    assert [spell["name"] for spell in updated.json()["spells"]] == ["魔能爆", "护盾术"]

    duplicate = await api_client.put(
        f"/api/v1/characters/{character_id}/spellbook",
        json={
            "revision": updated.json()["revision"],
            "spells": [updated.json()["spells"][0], updated.json()["spells"][0]],
        },
    )
    assert duplicate.status_code == 422


async def test_character_sheet_campaign_play_and_hp_flow(
    api_client: AsyncClient, tmp_path: Path
) -> None:
    character = await create_character(api_client, "赛蕾妮", max_hp=30)
    character_id = str(character["id"])
    await configure_sheet(api_client, tmp_path, character_id, "赛蕾妮")

    campaign_response = await api_client.post(
        "/api/v1/campaigns",
        json={"name": "鸦影之路", "description": "测试战役", "characterIds": [character_id]},
    )
    assert campaign_response.status_code == 201
    campaign = campaign_response.json()

    activation = await api_client.post(f"/api/v1/campaigns/{campaign['id']}:activate")
    assert activation.status_code == 200
    assert activation.json()["lifecycleStatus"] == "ACTIVE"

    play_response = await api_client.get(f"/api/v1/campaigns/{campaign['id']}/play")
    assert play_response.status_code == 200
    play_state = play_response.json()
    assert play_state["runtimeStatus"] == "IDLE"
    assert play_state["hpStates"][0]["currentHp"] == 30

    normal_check = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/skill-checks",
        json={
            "characterId": character_id,
            "skill": "perception",
            "rollMode": "NORMAL",
            "dc": 0,
            "clientRequestId": str(uuid4()),
        },
    )
    assert normal_check.status_code == 201
    assert normal_check.json()["dieTwo"] is None
    assert normal_check.json()["systemOutcome"] == "PASS"

    unresolved_check = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/skill-checks",
        json={
            "characterId": character_id,
            "skill": "investigation",
            "rollMode": "ADVANTAGE",
            "reason": "检查密室机关",
            "clientRequestId": str(uuid4()),
        },
    )
    assert unresolved_check.status_code == 201
    assert unresolved_check.json()["dieTwo"] is not None
    assert unresolved_check.json()["systemOutcome"] == "UNRESOLVED"
    adjudicated = await api_client.post(
        f"/api/v1/skill-checks/{unresolved_check.json()['id']}:adjudicate",
        json={"outcome": "PARTIAL_SUCCESS"},
    )
    assert adjudicated.status_code == 200
    assert adjudicated.json()["dmAdjudication"] == "PARTIAL_SUCCESS"

    duplicate_request_id = str(uuid4())
    first_roll = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/skill-checks",
        json={
            "characterId": character_id,
            "skill": "stealth",
            "rollMode": "DISADVANTAGE",
            "clientRequestId": duplicate_request_id,
        },
    )
    duplicate_roll = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/skill-checks",
        json={
            "characterId": character_id,
            "skill": "stealth",
            "rollMode": "DISADVANTAGE",
            "clientRequestId": duplicate_request_id,
        },
    )
    assert first_roll.status_code == duplicate_roll.status_code == 201
    assert first_roll.json()["id"] == duplicate_roll.json()["id"]

    request_id = str(uuid4())
    dm_message = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/messages",
        json={
            "content": "浓雾沿着石板路缓缓漫来。",
            "audience": "PUBLIC",
            "clientRequestId": request_id,
        },
    )
    assert dm_message.status_code == 201
    assert dm_message.json()["message"]["sequenceNo"] == 1
    assert dm_message.json()["message"]["recipients"][0]["characterId"] == character_id
    assert dm_message.json()["runtime"]["status"] == "IDLE"

    retried_message = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/messages",
        json={
            "content": "浓雾沿着石板路缓缓漫来。",
            "audience": "PUBLIC",
            "clientRequestId": request_id,
        },
    )
    assert retried_message.status_code == 201
    assert retried_message.json()["message"]["id"] == dm_message.json()["message"]["id"]

    stopped = await api_client.post(f"/api/v1/campaigns/{campaign['id']}/runtime:stop")
    assert stopped.status_code == 200
    assert stopped.json()["status"] == "IDLE"

    listed_messages = await api_client.get(f"/api/v1/campaigns/{campaign['id']}/messages")
    assert listed_messages.status_code == 200
    assert len(listed_messages.json()) == 1

    character_detail = await api_client.get(f"/api/v1/characters/{character_id}")
    revision = character_detail.json()["revision"]
    hp_reset = await api_client.patch(
        f"/api/v1/characters/{character_id}", json={"revision": revision, "maxHp": 42}
    )
    assert hp_reset.status_code == 200

    refreshed_play = await api_client.get(f"/api/v1/campaigns/{campaign['id']}/play")
    assert refreshed_play.json()["hpStates"][0]["currentHp"] == 42
    assert refreshed_play.json()["hpStates"][0]["maxHp"] == 42

    hp_update = await api_client.patch(
        f"/api/v1/campaigns/{campaign['id']}/characters/{character_id}/hp",
        json={"currentHp": 11},
    )
    assert hp_update.status_code == 200
    assert hp_update.json()["hpStates"][0]["currentHp"] == 11

    paused = await api_client.post(f"/api/v1/campaigns/{campaign['id']}:pause")
    assert paused.status_code == 200
    assert paused.json()["lifecycleStatus"] == "PAUSED"
    blocked_message = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/messages",
        json={
            "content": "暂停期间不应写入。",
            "audience": "PUBLIC",
            "clientRequestId": str(uuid4()),
        },
    )
    assert blocked_message.status_code == 409
    assert blocked_message.json()["errorCode"] == "CAMPAIGN_NOT_ACTIVE"

    resumed = await api_client.post(f"/api/v1/campaigns/{campaign['id']}:resume")
    assert resumed.status_code == 200
    assert resumed.json()["lifecycleStatus"] == "ACTIVE"
    resumed_play = await api_client.get(f"/api/v1/campaigns/{campaign['id']}/play")
    assert resumed_play.json()["hpStates"][0]["currentHp"] == 11

    completed = await api_client.post(f"/api/v1/campaigns/{campaign['id']}:complete")
    assert completed.status_code == 200
    assert completed.json()["lifecycleStatus"] == "COMPLETED"

    summaries = await api_client.get(f"/api/v1/campaigns/{campaign['id']}/summaries")
    assert summaries.status_code == 200
    assert {item["audience"] for item in summaries.json()} == {"DM", "CHARACTER"}

    exported = await api_client.get(f"/api/v1/campaigns/{campaign['id']}:export")
    assert exported.status_code == 200, exported.text
    export_payload = exported.json()
    assert export_payload["campaign"]["dmGuide"] == ""
    assert export_payload["campaign"]["sceneNotes"] == ""
    assert "skillChecks" in export_payload
    assert "dmDrafts" in export_payload


async def test_only_one_active_campaign(api_client: AsyncClient, tmp_path: Path) -> None:
    first = await create_character(api_client, "甲")
    second = await create_character(api_client, "乙")
    await configure_sheet(api_client, tmp_path, str(first["id"]), "甲")
    await configure_sheet(api_client, tmp_path, str(second["id"]), "乙")

    first_campaign = (
        await api_client.post(
            "/api/v1/campaigns",
            json={"name": "第一团", "characterIds": [first["id"]]},
        )
    ).json()
    second_campaign = (
        await api_client.post(
            "/api/v1/campaigns",
            json={"name": "第二团", "characterIds": [second["id"]]},
        )
    ).json()

    first_activation = await api_client.post(f"/api/v1/campaigns/{first_campaign['id']}:activate")
    second_activation = await api_client.post(f"/api/v1/campaigns/{second_campaign['id']}:activate")
    assert first_activation.status_code == 200
    assert second_activation.status_code == 409
    assert second_activation.json()["errorCode"] == "ACTIVE_CAMPAIGN_EXISTS"


async def test_incomplete_character_cannot_start_campaign(api_client: AsyncClient) -> None:
    character = await create_character(api_client, "未完成角色")
    campaign = (
        await api_client.post(
            "/api/v1/campaigns",
            json={"name": "未完成配置", "characterIds": [character["id"]]},
        )
    ).json()

    activation = await api_client.post(f"/api/v1/campaigns/{campaign['id']}:activate")
    assert activation.status_code == 400
    assert activation.json()["errorCode"] == "CHARACTER_CONFIGURATION_INCOMPLETE"


async def test_character_skill_set_requires_all_eighteen_skills(api_client: AsyncClient) -> None:
    character = await create_character(api_client, "技能角色")
    skills = await api_client.get(f"/api/v1/characters/{character['id']}/skills")
    assert skills.status_code == 200
    assert len(skills.json()["modifiers"]) == 18
    assert set(skills.json()["modifiers"].values()) == {0}

    updated = dict(skills.json()["modifiers"])
    updated["investigation"] = 5
    saved = await api_client.put(
        f"/api/v1/characters/{character['id']}/skills",
        json={"revision": skills.json()["revision"], "modifiers": updated},
    )
    assert saved.status_code == 200
    assert saved.json()["modifiers"]["investigation"] == 5

    invalid = await api_client.put(
        f"/api/v1/characters/{character['id']}/skills",
        json={"revision": saved.json()["revision"], "modifiers": {"arcana": 2}},
    )
    assert invalid.status_code == 422


async def test_campaign_module_can_create_opening_draft(
    api_client: AsyncClient, monkeypatch
) -> None:
    monkeypatch.setattr(
        "backend.app.services.dm_draft_service.get_settings",
        lambda: Settings(character_model="", deepseek_api_key=None),
    )
    campaign = (
        await api_client.post(
            "/api/v1/campaigns",
            json={
                "name": "开场草稿测试",
                "moduleContent": "古堡在暴雨中重新亮起灯火。",
                "styleInstructions": "克制、悬疑。",
                "openingInstructions": "从角色抵达古堡开始。",
            },
        )
    ).json()
    generated = await api_client.post(
        f"/api/v1/campaigns/{campaign['id']}/dm-drafts:generate-opening"
    )
    assert generated.status_code == 200
    assert generated.json()["triggerType"] == "OPENING"
    assert generated.json()["status"] in {"READY", "FAILED"}
    assert generated.json()["sessionId"] is None


async def test_character_acquaintance_cannot_be_manually_edited(
    api_client: AsyncClient,
) -> None:
    first = await create_character(api_client, "全局相识甲")
    second = await create_character(api_client, "全局相识乙")
    member_ids = [first["id"], second["id"]]
    first_campaign = (
        await api_client.post(
            "/api/v1/campaigns",
            json={"name": "相识来源团", "characterIds": member_ids},
        )
    ).json()
    second_campaign = (
        await api_client.post(
            "/api/v1/campaigns",
            json={"name": "相识继承团", "characterIds": member_ids},
        )
    ).json()

    initial = await api_client.get(
        f"/api/v1/campaigns/{first_campaign['id']}/acquaintances"
    )
    assert initial.status_code == 200
    assert initial.json()[0]["acquainted"] is False

    manual_update = await api_client.put(
        f"/api/v1/campaigns/{first_campaign['id']}/acquaintances",
        json={
            "characterAId": first["id"],
            "characterBId": second["id"],
            "acquainted": True,
        },
    )
    assert manual_update.status_code == 405

    inherited = await api_client.get(
        f"/api/v1/campaigns/{second_campaign['id']}/acquaintances"
    )
    assert inherited.status_code == 200
    assert inherited.json()[0]["acquainted"] is False
