from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, Field

from backend.app.core.errors import AppError
from backend.app.domain.enums import SkillName

PARSER_VERSION = "sad-spirit-v1"
REQUIRED_SHEETS = {"主要", "起源", "背包", "法术大全"}

SKILL_LABELS: dict[str, SkillName] = {
    "运动": SkillName.ATHLETICS,
    "特技": SkillName.ACROBATICS,
    "巧手": SkillName.SLEIGHT_OF_HAND,
    "隐匿": SkillName.STEALTH,
    "调查": SkillName.INVESTIGATION,
    "奥秘": SkillName.ARCANA,
    "历史": SkillName.HISTORY,
    "自然": SkillName.NATURE,
    "宗教": SkillName.RELIGION,
    "察觉": SkillName.PERCEPTION,
    "洞悉": SkillName.INSIGHT,
    "驯兽": SkillName.ANIMAL_HANDLING,
    "医药": SkillName.MEDICINE,
    "医疗": SkillName.MEDICINE,
    "求生": SkillName.SURVIVAL,
    "游说": SkillName.PERSUASION,
    "欺瞒": SkillName.DECEPTION,
    "威吓": SkillName.INTIMIDATION,
    "表演": SkillName.PERFORMANCE,
}
SKILL_ROW_RANGE = (28, 58)


def _to_camel(value: str) -> str:
    first, *rest = value.split("_")
    return first + "".join(part.capitalize() for part in rest)


class SnapshotModel(BaseModel):
    model_config = ConfigDict(alias_generator=_to_camel, populate_by_name=True)


class NamedDescription(SnapshotModel):
    name: str
    description: str = ""


class SpellSnapshot(NamedDescription):
    category: str
    school: str = ""
    casting_time: str = ""
    range: str = ""
    duration: str = ""
    components: str = ""


def _named_description_list() -> list[NamedDescription]:
    return []


def _spell_list() -> list[SpellSnapshot]:
    return []


class CharacterSheetSnapshot(SnapshotModel):
    parser_version: str = PARSER_VERSION
    character_name: str
    species: str
    subspecies: str | None = None
    classes: list[str]
    subclasses: list[str]
    background: str
    homeland: str | None = None
    languages: list[str] = Field(default_factory=list)
    tools: list[str] = Field(default_factory=list)
    class_features: list[NamedDescription] = Field(default_factory=_named_description_list)
    species_features: list[NamedDescription] = Field(default_factory=_named_description_list)
    feats: list[NamedDescription] = Field(default_factory=_named_description_list)
    spells: list[SpellSnapshot] = Field(default_factory=_spell_list)
    equipment: list[NamedDescription] = Field(default_factory=_named_description_list)
    inventory: list[NamedDescription] = Field(default_factory=_named_description_list)
    skills: dict[SkillName, int] = Field(default_factory=dict)


def _text(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    text = str(value).strip()
    if text in {"#NAME?", "#N/A", "#REF!", "-", "—", ">", "<", "O"}:
        return ""
    return text


def _cell_text(sheet: Worksheet, address: str) -> str:
    return _text(sheet[address].value)


def _text_cells(sheet: Worksheet, cell_range: str, *, excluded: set[str]) -> list[str]:
    values: list[str] = []
    for row in sheet[cell_range]:
        for cell in row:
            value = _text(cell.value)
            if value and value not in excluded:
                values.append(value)
    return list(dict.fromkeys(values))


def _feature_rows(
    sheet: Worksheet, name_column: str, description_column: str, start: int, end: int
) -> list[NamedDescription]:
    features: list[NamedDescription] = []
    for row in range(start, end + 1):
        name = _cell_text(sheet, f"{name_column}{row}")
        description = _cell_text(sheet, f"{description_column}{row}")
        if name and name not in {"名称", "描述", "Lv"}:
            features.append(NamedDescription(name=name, description=description))
    return features


def _spell_catalog(sheet: Worksheet) -> dict[str, dict[str, str]]:
    catalog: dict[str, dict[str, str]] = {}
    for row in sheet.iter_rows(min_row=3, max_col=13):
        name = _text(row[0].value)
        if not name:
            continue
        components = "".join(
            label
            for label, cell in zip(("V", "S", "M"), row[7:10], strict=True)
            if _text(cell.value)
        )
        catalog[name] = {
            "school": _text(row[2].value),
            "casting_time": _text(row[5].value),
            "range": _text(row[6].value),
            "components": components,
            "duration": _text(row[11].value),
            "description": _text(row[12].value),
        }
    return catalog


def _selected_spells(main: Worksheet, catalog: dict[str, dict[str, str]]) -> list[SpellSnapshot]:
    selections = {
        "AT_WILL": (("Q", 57, 64),),
        "PREPARED": (("Y", 57, 64), ("AG", 57, 64), ("AO", 57, 64)),
        "CANTRIP": (("Q", 68, 72), ("Y", 68, 72)),
    }
    spells: list[SpellSnapshot] = []
    seen: set[tuple[str, str]] = set()
    for category, ranges in selections.items():
        for column, start, end in ranges:
            for row in range(start, end + 1):
                name = _cell_text(main, f"{column}{row}")
                key = (category, name)
                if not name or name.isdigit() or key in seen:
                    continue
                seen.add(key)
                details = catalog.get(name, {})
                spells.append(SpellSnapshot(name=name, category=category, **details))
    return spells


def _equipment(main: Worksheet) -> list[NamedDescription]:
    items: list[NamedDescription] = []
    armor_name = _cell_text(main, "L31")
    if armor_name:
        items.append(NamedDescription(name=armor_name, description=_cell_text(main, "V31")))
    for row in range(33, 38):
        name = _cell_text(main, f"L{row}")
        if name:
            items.append(NamedDescription(name=name, description=_cell_text(main, f"V{row}")))
    for row in range(39, 49):
        name = _cell_text(main, f"L{row}")
        if name:
            items.append(NamedDescription(name=name, description=_cell_text(main, f"T{row}")))
    return items


def _skill_modifier(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    try:
        return int(str(value).strip())
    except ValueError:
        return None


def _skills(main: Worksheet) -> dict[SkillName, int]:
    start, end = SKILL_ROW_RANGE
    skills: dict[SkillName, int] = {}
    for row in range(start, end + 1):
        skill = SKILL_LABELS.get(_cell_text(main, f"C{row}"))
        if skill is None or skill in skills:
            continue
        modifier = _skill_modifier(main[f"I{row}"].value)
        if modifier is not None:
            skills[skill] = modifier
    return skills


def _inventory(backpack: Worksheet) -> list[NamedDescription]:
    items: list[NamedDescription] = []
    for start, end in ((5, 14), (16, 25), (29, 38), (40, 49)):
        for row in range(start, end + 1):
            name = _cell_text(backpack, f"Y{row}")
            if name:
                items.append(
                    NamedDescription(name=name, description=_cell_text(backpack, f"AC{row}"))
                )
    return items


class CharacterSheetParser:
    def parse(self, path: Path) -> CharacterSheetSnapshot:
        try:
            workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
        except Exception as error:
            raise AppError("SHEET_OPEN_FAILED", "无法读取角色卡。") from error

        try:
            missing = REQUIRED_SHEETS.difference(workbook.sheetnames)
            if missing:
                raise AppError(
                    "SHEET_TEMPLATE_MISMATCH",
                    "角色卡缺少模板要求的工作表。",
                    details={"missingSheets": sorted(missing)},
                )
            main = workbook["主要"]
            origin = workbook["起源"]
            backpack = workbook["背包"]
            spell_database = workbook["法术大全"]
            if not _cell_text(main, "A1").startswith("DND 5E2024 人物卡<悲灵"):
                raise AppError("SHEET_TEMPLATE_MISMATCH", "角色卡不是受支持的固定模板。")

            classes = [_cell_text(main, address) for address in ("E6", "E7", "E8")]
            subclasses = [_cell_text(main, address) for address in ("I6", "I7", "I8")]
            snapshot = CharacterSheetSnapshot(
                character_name=_cell_text(main, "E3"),
                species=_cell_text(main, "T6"),
                subspecies=_cell_text(main, "T7") or None,
                classes=[value for value in classes if value],
                subclasses=[value for value in subclasses if value],
                background=_cell_text(origin, "E6"),
                homeland=_cell_text(origin, "E5") or None,
                languages=_text_cells(origin, "B24:M28", excluded={"语言"}),
                tools=_text_cells(origin, "B31:M31", excluded={"熟练工具"}),
                class_features=_feature_rows(main, "AX", "BC", 17, 40),
                species_features=_feature_rows(main, "BT", "BZ", 30, 40),
                feats=_feature_rows(main, "BU", "BZ", 43, 54),
                spells=_selected_spells(main, _spell_catalog(spell_database)),
                equipment=_equipment(main),
                inventory=_inventory(backpack),
                skills=_skills(main),
            )
            if not snapshot.character_name or not snapshot.species or not snapshot.classes:
                raise AppError("SHEET_REQUIRED_DATA_MISSING", "角色卡缺少姓名、种族或职业。")
            return snapshot
        finally:
            workbook.close()
